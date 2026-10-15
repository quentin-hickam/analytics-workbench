# Kept cases:
# test_renderer_builds_every_slide_verbatim: the reference renderer builds the title slide and one slide per outline slide in order, with headlines, bullets and nested bullets, and notes verbatim, and leaves the draft untouched.
# test_renderer_starts_the_appendix_with_a_divider: a draft with ## Appendix gets one divider slide titled Appendix, holding nothing else, before its slides, which keep their order.
# test_renderer_draws_every_form_from_its_file: each of the seven forms is drawn from its chart file at the file's values and order: bar, column, line, and stacked types with a series pivoted, intervals as custom error bars, each bar, column, and line value axis from zero, scatter points and fit, a table's cells; alt text and the source line (with the **Interval:** text) and caveat line come from the specification.
# test_renderer_line_axis_follows_axis_field: a line chart's value axis includes zero unless **Axis:** says non-zero baseline, which then joins the source line; other wording, the template's "reason for another" included, keeps zero, and so does a column chart whose **Axis:** says non-zero baseline, leaving it out of the source line.
# test_renderer_value_axis_spans_zero_and_every_value: a bar, column, or line axis pins zero as its minimum when every value and interval bound is zero or more, and as its maximum when every one is zero or less and one is below zero; it pins an interval bound that reaches past every value, whichever side of zero it lies on; otherwise it leaves the bound automatic, so mixed-sign and all-negative values and every interval stay visible.
# test_renderer_is_deterministic: two renders of one draft are byte-identical.
# test_renderer_refuses_a_draft_breaking_the_contract: an outline or chart row stops the render, naming the rows, and writes nothing.
# test_renderer_refuses_data_it_cannot_draw: a repeated category and series pair, a blank interval cell, or a negative part of a stacked bar stops the render with a ContractError instead of dropping a row or crashing.
# The python-pptx tests skip when the package is missing; the renderer is a test fixture, never shipped.

import importlib.util
from pathlib import Path

import pytest

pytest.importorskip("pptx")

_spec = importlib.util.spec_from_file_location("deck_renderer", Path(__file__).with_name("deck_renderer.py"))
renderer = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(renderer)

from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

FORMS = [
    ("Horizontal bar", "category: Region; value: Closed late (%)", "Region,Closed late (%)\nWest,21.3\nEast,2.7\nNorth,8\n"),
    ("Bar with intervals", "category: Region; value: Rate (%); lower: Lower (%); upper: Upper (%)",
     "Region,Rate (%),Lower (%),Upper (%)\nWest,21,18,25\nEast,5,4,6.5\n"),
    ("Line", "category: Month; series: Region; value: Cases", "Month,Region,Cases\nJan,West,12\nJan,East,9\n"
                                                              "Feb,West,15\nFeb,East,8\nMar,West,18\n"),
    ("Column", "category: Month; value: Cases; value: Late", "Month,Cases,Late\nJan,1204,96\nFeb,1310,101\n"),
    ("Stacked bar", "category: Region; value: Late; value: On time", "Region,Late,On time\nWest,21,79\nEast,5,95\n"),
    ("Scatter", "x: Spend; y: Late (%); label: Store; fit: Fitted (%)",
     "Store,Spend,Late (%),Fitted (%)\nA,10,4,4.5\nB,20,9,8.5\nC,30,12,12.5\n"),
    ("Table", "column: Region; column: Cases", "Region,Cases\nWest,1240\nEast,980\n"),
]


def findings():
    slides = "".join(f"""### Finding {n} shows the {form.lower()}

- Point {n}.
- **Caveat:** caveat {n}.

**Chart {n}.** {form}: comparison {n}
- **Data:** {data}
- **Highlight:** none
- **Axis:** zero baseline
- **Interval:** {'bars show 95% intervals' if form == 'Bar with intervals' else 'none'}
- **Source:** source {n}
- **Caveat:** {'none' if n % 2 else f'value {n} is partial'}
- **Alt text:** alt text {n}

**Notes:** notes {n}

""" for n, (form, data, _) in enumerate(FORMS, start=1))
    return f"""# Late claims in the west
Regional operations managers, claims closed in the first half of the year

## Answer

**Headline:** West claims close late four times as often as the east

- The west closed most claims late.
  - Nested under the first bullet.
- Staffing is the decision.

**Notes:** The answer in brief.

## Findings

{slides}## How we know

**Headline:** Totals reconcile to the finance ledger

- Claims are counted once.

**Notes:** The ledger is the finance team's.

## Limitations and open caveats

**Headline:** March records are incomplete in every region

- **Caveat:** March counts may rise.

**Notes:** The full March file comes in July.
"""


@pytest.fixture
def package(tmp_path):
    directory = tmp_path / "draft"
    (directory / "charts").mkdir(parents=True)
    (directory / "findings.md").write_text(findings(), encoding="utf-8")
    for n, (_, _, content) in enumerate(FORMS, start=1):
        (directory / f"charts/chart-{n}.csv").write_text(content, encoding="utf-8")
    return directory


def deck(package, tmp_path):
    return Presentation(renderer.render(package, tmp_path / "deck.pptx"))


def body(slide):
    return [(p.level, p.text) for p in slide.placeholders[1].text_frame.paragraphs]


def test_renderer_builds_every_slide_verbatim(package, tmp_path):
    before = {path: path.read_bytes() for path in package.rglob("*") if path.is_file()}
    slides = list(deck(package, tmp_path).slides)
    assert {path: path.read_bytes() for path in package.rglob("*") if path.is_file()} == before
    assert [slide.shapes.title.text for slide in slides] == [
        "Late claims in the west", "West claims close late four times as often as the east",
        *(f"Finding {n} shows the {form.lower()}" for n, (form, _, _) in enumerate(FORMS, start=1)),
        "Totals reconcile to the finance ledger", "March records are incomplete in every region"]
    assert slides[0].placeholders[1].text == ("Regional operations managers, claims closed in the first half of "
                                              "the year")
    assert body(slides[1]) == [(0, "The west closed most claims late."), (1, "Nested under the first bullet."),
                               (0, "Staffing is the decision.")]
    assert body(slides[2]) == [(0, "Point 1."), (0, "Caveat: caveat 1.")]
    assert [slide.notes_slide.notes_text_frame.text for slide in slides[1:]] == [
        "The answer in brief.", *(f"notes {n}" for n in range(1, 8)), "The ledger is the finance team's.",
        "The full March file comes in July."]
    assert not slides[0].has_notes_slide


APPENDIX = """
## Appendix

### Late claims fell in the north

- North closed 4% of claims late.

**Notes:** The north is the smallest region.

### How we know

**Headline:** Totals reconcile to the finance ledger

- Claims are counted once.

**Notes:** The ledger is the finance team's.
"""


def test_renderer_starts_the_appendix_with_a_divider(package, tmp_path):
    path = package / "findings.md"
    text = path.read_text(encoding="utf-8")
    how = text[text.index("## How we know"):text.index("## Limitations")]
    path.write_text(text.replace(how, "") + APPENDIX, encoding="utf-8")
    slides = list(deck(package, tmp_path).slides)
    titles = [slide.shapes.title.text for slide in slides]
    assert titles[-4:] == ["March records are incomplete in every region", "Appendix", "Late claims fell in the north",
                           "Totals reconcile to the finance ledger"]
    divider = slides[-3]
    assert [shape.text_frame.text for shape in divider.shapes if shape.has_text_frame] == ["Appendix"]
    assert not divider.has_notes_slide
    assert body(slides[-2]) == [(0, "North closed 4% of claims late.")]


def frames(slide):
    return [shape for shape in slide.shapes if shape.has_chart or shape.has_table]


def test_renderer_draws_every_form_from_its_file(package, tmp_path):
    slides = list(deck(package, tmp_path).slides)[2:9]
    charts = [frames(slide)[0] for slide in slides]
    assert [frame.chart.chart_type for frame in charts[:6]] == [
        XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE_MARKERS,
        XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_STACKED_100, XL_CHART_TYPE.XY_SCATTER]

    def series(frame):
        return [(s.name, list(s.values)) for s in frame.chart.plots[0].series]

    bar, intervals, line, column, stacked, scatter = (frame for frame in charts[:6])
    assert list(bar.chart.plots[0].categories) == ["West", "East", "North"]
    assert series(bar) == [("Closed late (%)", [21.3, 2.7, 8.0])]
    assert bar.chart.category_axis.reverse_order and bar.chart.value_axis.minimum_scale == 0
    assert series(intervals) == [("Rate (%)", [21.0, 5.0])]
    bars = intervals.chart._chartSpace.xpath(".//c:errBars")[0]
    assert [[float(v.text) for v in bars.find(qn(side)).iter(qn("c:v"))] for side in ("c:plus", "c:minus")] == [
        [4.0, 1.5], [3.0, 1.0]]
    assert list(line.chart.plots[0].categories) == ["Jan", "Feb", "Mar"]
    assert series(line) == [("West", [12.0, 15.0, 18.0]), ("East", [9.0, 8.0, None])]
    assert line.chart.value_axis.minimum_scale == 0
    assert series(column) == [("Cases", [1204.0, 1310.0]), ("Late", [96.0, 101.0])]
    assert series(stacked) == [("Late", [21.0, 5.0]), ("On time", [79.0, 95.0])]
    assert series(scatter) == [("Late (%)", [4.0, 9.0, 12.0]), ("Fitted (%)", [4.5, 8.5, 12.5])]
    assert [float(v.text) for v in scatter.chart._chartSpace.xpath(".//c:ser")[0].find(qn("c:xVal")).iter(
        qn("c:v"))] == [10.0, 20.0, 30.0]
    assert [p.data_label.text_frame.text for p in scatter.chart.plots[0].series[0].points] == ["A", "B", "C"]
    table = charts[6].table
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["Region", "Cases"], ["West", "1240"], ["East", "980"]]
    for n, (slide, frame) in enumerate(zip(slides, charts), start=1):
        assert frame._element.xpath("./*[1]/p:cNvPr")[0].get("descr") == f"alt text {n}"
        source = [shape.text_frame.text for shape in slide.shapes
                  if shape.has_text_frame and shape.text_frame.text.startswith("source")]
        expected = f"source {n}; bars show 95% intervals" if n == 2 else f"source {n}"
        assert source == [expected if n % 2 else f"{expected}\n* value {n} is partial"]


def source_line(slide):
    return [shape.text_frame.text for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.startswith("source")]


@pytest.mark.parametrize("number, axis, minimum, source", [
    (3, "non-zero baseline, to show the small monthly change", None,
     "source 3; non-zero baseline, to show the small monthly change"),
    (3, "Nonzero baseline: monthly changes are small", None, "source 3; Nonzero baseline: monthly changes are small"),
    (3, "cases, starting at 8 to show the small monthly change", 0, "source 3"),
    (3, "non-zero", 0, "source 3"),
    (4, "non-zero baseline, because the counts are large", 0, "source 4\n* value 4 is partial"),
])
def test_renderer_line_axis_follows_axis_field(package, tmp_path, number, axis, minimum, source):
    findings_path = package / "findings.md"
    text = findings_path.read_text(encoding="utf-8")
    line = text.index(f"**Chart {number}.**")
    findings_path.write_text(text[:line] + text[line:].replace("zero baseline", axis, 1), encoding="utf-8")
    slide = list(deck(package, tmp_path).slides)[number + 1]
    frame = frames(slide)[0]
    assert frame.chart.chart_type == (XL_CHART_TYPE.LINE_MARKERS if number == 3 else XL_CHART_TYPE.COLUMN_CLUSTERED)
    assert (frame.chart.value_axis.minimum_scale, frame.chart.value_axis.maximum_scale) == (minimum, None)
    assert source_line(slide) == [source]


@pytest.mark.parametrize("number, content, minimum, maximum", [
    (3, "Month,Region,Cases\nJan,West,-4.3\nFeb,West,2\nJan,East,1\n", None, None),
    (3, "Month,Region,Cases\nJan,West,-4\nFeb,West,-6\nJan,East,-1\n", None, 0),
    (3, "Month,Region,Cases\nJan,West,0\nFeb,West,0\n", 0, None),
    (1, "Region,Closed late (%)\nWest,-2.5\nEast,-0.5\nNorth,-8\n", None, 0),
    (4, "Month,Cases,Late\nJan,-3,1\nFeb,2,-5\n", None, None),
    (2, "Region,Rate (%),Lower (%),Upper (%)\nWest,2,-1,5\nEast,5,4,6.5\n", -1, 6.5),
    (2, "Region,Rate (%),Lower (%),Upper (%)\nWest,2,1,3\nEast,5,4,4.5\n", 0, None),
    (2, "Region,Rate (%),Lower (%),Upper (%)\nWest,-2,-3,-0.5\nEast,-5,-7,-4\n", -7, 0),
    (2, "Region,Rate (%),Lower (%),Upper (%)\nWest,-2,-3,1\nEast,-5,-7,-4\n", -7, 1),
    (2, "Region,Rate (%),Lower (%),Upper (%)\nWest,-2,-3,-1\nEast,3,2,4\n", -3, 4),
])
def test_renderer_value_axis_spans_zero_and_every_value(package, tmp_path, number, content, minimum, maximum):
    (package / f"charts/chart-{number}.csv").write_text(content, encoding="utf-8")
    axis = frames(list(deck(package, tmp_path).slides)[number + 1])[0].chart.value_axis
    assert (axis.minimum_scale, axis.maximum_scale) == (minimum, maximum)


def test_renderer_is_deterministic(package, tmp_path):
    first = renderer.render(package, tmp_path / "first.pptx").read_bytes()
    assert renderer.render(package, tmp_path / "second.pptx").read_bytes() == first


@pytest.mark.parametrize("old, new, row", [
    ("**Headline:** Totals reconcile to the finance ledger",
     "**Headline:** Totals reconcile to the finance ledger and to every regional team's own records in each month",
     {"slide": "How we know", "problem": "the headline has 16 words; keep it to 15"}),
    ("Column: comparison 4", "Pie: comparison 4",
     {"chart": "Chart 4", "problem": "names no chart form (Pie); open the specification with one of: Horizontal "
                                     "bar, Bar with intervals, Line, Column, Stacked bar, Scatter, Table"}),
])
def test_renderer_refuses_a_draft_breaking_the_contract(package, tmp_path, old, new, row):
    path = package / "findings.md"
    path.write_text(path.read_text().replace(old, new))
    with pytest.raises(renderer.ContractError) as refused:
        renderer.render(package, tmp_path / "deck.pptx")
    assert [{key: value for key, value in found.items() if key != "line"} for found in refused.value.rows] == [row]
    assert not (tmp_path / "deck.pptx").exists()


@pytest.mark.parametrize("number, content, problem", [
    (3, "Month,Region,Cases\nJan,West,12\nJan,West,99\nFeb,West,15\n",
     "'Jan' and 'West' repeat a row, so one point would hide the other"),
    (2, "Region,Rate (%),Lower (%),Upper (%)\nWest,21,,25\nEast,5,4,6.5\n",
     "a bar's value, lower, or upper cell is blank; each interval needs all three"),
    (5, "Region,Late,On time\nWest,-21,79\nEast,5,95\n",
     "value -21 is below zero, which a Stacked bar cannot show as a part of a whole"),
])
def test_renderer_refuses_data_it_cannot_draw(package, tmp_path, number, content, problem):
    (package / f"charts/chart-{number}.csv").write_text(content, encoding="utf-8")
    with pytest.raises(renderer.ContractError) as refused:
        renderer.render(package, tmp_path / "deck.pptx")
    assert refused.value.rows == [{"chart": f"Chart {number}", "problem": problem}]
    assert not (tmp_path / "deck.pptx").exists()
