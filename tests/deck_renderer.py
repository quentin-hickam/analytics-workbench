"""Reference renderer: build a .pptx from a package draft's findings.md and charts/, deterministically.

A test fixture, kept out of the distributed archive. It proves the deck contract in
package-format/deck-rules.md is complete enough to build a deck from without Microsoft 365: everything it
places comes from the outline `awb_draft.parse_outline` derives and from the chart files. It follows the
content rules (one slide per outline slide, headlines, bullets, and notes verbatim, every chart drawn from
its file at the file's values, each bar, column, and line value axis spanning zero and every value and
interval bound unless a line chart's **Axis:** says non-zero baseline, a Stacked bar's parts zero or more,
intervals as error bars, the source line with the baseline's reason, what the
intervals mean, and the caveat's reason, and alt text from the specification) and leaves styling, sections,
and theme to the builders users run. It marks no value with the caveat's asterisk: **Caveat:** names that
value only in words. Like any builder that cannot ask, it
refuses a draft that `check-draft`'s outline and chart checks fail, naming each row.

    python3 tests/deck_renderer.py <draft directory> <output.pptx>
"""

import csv
import importlib.util
import io
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

_ASSET = Path(__file__).resolve().parents[1] / ".agents/skills/awb-package/assets/awb_draft.py"
_spec = importlib.util.spec_from_file_location("awb_draft_renderer", _ASSET)
draft = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(draft)

CHART_TYPES = {
    "Horizontal bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "Bar with intervals": XL_CHART_TYPE.BAR_CLUSTERED,
    "Line": XL_CHART_TYPE.LINE_MARKERS,
    "Column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "Stacked bar": XL_CHART_TYPE.BAR_STACKED_100,
    "Scatter": XL_CHART_TYPE.XY_SCATTER,
}
_FIXED_TIME = datetime(2000, 1, 1)
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_TITLE, _CONTENT = 0, 1  # the default template's layouts
_CHART_BOX = (Inches(5.0), Inches(1.6), Inches(4.6), Inches(4.6))
_NON_ZERO = re.compile(r"\bnon-?zero baseline\b", re.IGNORECASE)


class ContractError(ValueError):
    """The draft breaks the deck contract; .rows holds each outline and chart row."""

    def __init__(self, rows):
        super().__init__("; ".join(f"{row.get('slide') or row.get('chart')}: {row['problem']}" for row in rows))
        self.rows = rows


def read_chart(path):
    """A chart file as (headers, rows of text), in the file's order."""
    with Path(path).open(encoding="utf-8", newline="") as stream:
        lines = [row for row in csv.reader(stream) if row]
    return lines[0], lines[1:]


def _number(text, where):
    if text.strip() == "":
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        raise ContractError([{"chart": where, "problem": f"value {text!r} is not a number"}]) from None


def _required(text, where):
    """A number the chart cannot draw without; a blank cell is refused."""
    if (number := _number(text, where)) is None:
        raise ContractError([{"chart": where, "problem": "a bar's value, lower, or upper cell is blank; "
                                                    "each interval needs all three"}])
    return number


def _roles(pairs, role):
    return [header for name, header in pairs if name == role]


def category_data(chart, headers, rows):
    """(categories, [(series name, values)]) for a category form: every row in the file's order, or with a
    series column, each category and series in the order it first appears."""
    pairs, _ = draft.chart_data(chart)
    where = f"Chart {chart['number']}"
    column = {header: index for index, header in enumerate(headers)}
    category = column[_roles(pairs, "category")[0]]
    values = _roles(pairs, "value")
    if _roles(pairs, "series"):
        categories = list(dict.fromkeys(row[category] for row in rows))
        series, value = column[_roles(pairs, "series")[0]], column[values[0]]
        cells = {}
        for row in rows:
            if (row[category], row[series]) in cells:
                raise ContractError([{"chart": where, "problem": f"{row[category]!r} and {row[series]!r} repeat a "
                                                                 "row, so one point would hide the other"}])
            cells[row[category], row[series]] = _number(row[value], where)
        names = list(dict.fromkeys(row[series] for row in rows))
        return categories, [(name, [cells.get((c, name)) for c in categories]) for name in names]
    return [row[category] for row in rows], [(name, [_number(row[column[name]], where) for row in rows])
                                             for name in values]


def _error_bars(chart, pairs, headers, rows, where):
    """Custom error bars from the file's lower and upper columns, on the one bar series."""
    column = {header: index for index, header in enumerate(headers)}
    value, lower, upper = (column[_roles(pairs, role)[0]] for role in ("value", "lower", "upper"))
    plus = [_required(row[upper], where) - _required(row[value], where) for row in rows]
    minus = [_required(row[value], where) - _required(row[lower], where) for row in rows]
    ser = chart._chartSpace.xpath(".//c:ser")[0]
    bars = etree.SubElement(ser, qn("c:errBars"))
    for tag, val in (("c:errBarType", "both"), ("c:errValType", "cust"), ("c:noEndCap", "0")):
        etree.SubElement(bars, qn(tag)).set("val", val)
    for tag, numbers in (("c:plus", plus), ("c:minus", minus)):
        literal = etree.SubElement(etree.SubElement(bars, qn(tag)), qn("c:numLit"))
        etree.SubElement(literal, qn("c:formatCode")).text = "General"
        etree.SubElement(literal, qn("c:ptCount")).set("val", str(len(numbers)))
        for index, number in enumerate(numbers):
            point = etree.SubElement(literal, qn("c:pt"))
            point.set("idx", str(index))
            etree.SubElement(point, qn("c:v")).text = repr(number)
    ser.remove(bars)  # c:errBars sits before c:cat in a bar series
    ser.find(qn("c:cat")).addprevious(bars)


def non_zero_baseline(chart):
    """Whether a Line specification's **Axis:** says non-zero baseline (or nonzero, in any case); any other
    wording, and any other form, keeps zero on the value axis."""
    return chart["form"] == "Line" and bool(_NON_ZERO.search(chart["fields"].get("Axis", "")))


def _place_chart(slide, chart, draft_dir):
    """Draw a chart specification from its file; return the graphic frame."""
    headers, rows = read_chart(Path(draft_dir) / "charts" / f"chart-{chart['number']}.csv")
    pairs, _ = draft.chart_data(chart)
    where, form = f"Chart {chart['number']}", chart["form"]
    if form == "Table":
        columns = [headers.index(header) for header in _roles(pairs, "column")]
        frame = slide.shapes.add_table(len(rows) + 1, len(columns), *_CHART_BOX)
        for position, index in enumerate(columns):
            frame.table.cell(0, position).text = headers[index]
            for number, row in enumerate(rows, start=1):
                frame.table.cell(number, position).text = row[index]
        return frame
    if form == "Scatter":
        column = {header: index for index, header in enumerate(headers)}
        x, y = column[_roles(pairs, "x")[0]], column[_roles(pairs, "y")[0]]
        data = XyChartData()
        points = data.add_series(headers[y])
        for row in rows:
            points.add_data_point(_number(row[x], where), _number(row[y], where))
        if _roles(pairs, "fit"):
            fit = column[_roles(pairs, "fit")[0]]
            fitted = data.add_series(headers[fit])
            for row in rows:
                fitted.add_data_point(_number(row[x], where), _number(row[fit], where))
        frame = slide.shapes.add_chart(CHART_TYPES[form], *_CHART_BOX, data)
        if _roles(pairs, "label"):
            label = column[_roles(pairs, "label")[0]]
            for point, row in zip(frame.chart.plots[0].series[0].points, rows):
                point.data_label.text_frame.text = row[label]
        return frame
    categories, series = category_data(chart, headers, rows)
    plotted = [value for _, values in series for value in values if value is not None]
    if form == "Stacked bar" and min(plotted, default=0) < 0:
        raise ContractError([{"chart": where, "problem": f"value {min(plotted):g} is below zero, which a Stacked bar "
                                                         "cannot show as a part of a whole"}])
    if form == "Bar with intervals":
        column = {header: index for index, header in enumerate(headers)}
        plotted += [number for role in ("lower", "upper") for row in rows
                    if (number := _number(row[column[_roles(pairs, role)[0]]], where)) is not None]
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    frame = slide.shapes.add_chart(CHART_TYPES[form], *_CHART_BOX, data)
    if form in ("Horizontal bar", "Bar with intervals", "Stacked bar"):
        frame.chart.category_axis.reverse_order = True  # the file's first row on top
    # The value axis spans zero and every plotted value and interval bound, so a negative value stays in view;
    # a line chart leaves the bounds to the axis when **Axis:** says non-zero baseline.
    if not non_zero_baseline(chart):
        frame.chart.value_axis.minimum_scale = min([0, *plotted])
        if max(plotted, default=0) < 0:
            frame.chart.value_axis.maximum_scale = 0
    if form == "Bar with intervals":
        _error_bars(frame.chart, pairs, headers, rows, where)
    frame.chart.has_legend = len(series) > 1
    return frame


def _text(frame, lines):
    """Fill a text frame with (level, text) lines."""
    frame.clear()
    for index, (level, text) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text, paragraph.level = text, level


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _normalize(blob):
    """The zip archive with fixed member times, and fixed creation times in any embedded workbook."""
    source = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.endswith(".xlsx"):
                data = _normalize(data)
            elif info.filename.endswith("core.xml"):
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2000-01-01T00:00:00Z", data)
            target.writestr(zipfile.ZipInfo(info.filename, _ZIP_TIME), data, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def render(draft_dir, output):
    """Build the deck for the draft at draft_dir into output; refuse a draft that breaks the contract."""
    draft_dir, output = Path(draft_dir), Path(output)
    findings = draft_dir / "findings.md"
    rows = draft.outline_problems(findings) + draft.check_charts(findings, draft_dir / "charts")
    if rows:
        raise ContractError(rows)
    outline = draft.read_outline(findings)
    deck = Presentation()
    properties = deck.core_properties
    properties.title, properties.author, properties.last_modified_by = outline["title"]["text"], "", ""
    properties.created = properties.modified = properties.last_printed = _FIXED_TIME
    properties.revision = 1

    title = deck.slides.add_slide(deck.slide_layouts[_TITLE])
    title.shapes.title.text = outline["title"]["text"]
    title.placeholders[1].text = outline["title"]["subtitle"]
    _notes(title, outline["title"]["notes"])
    charts = {chart["number"]: chart for chart in outline["charts"]}
    for spec in outline["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[_CONTENT])
        slide.shapes.title.text = spec["headline"]
        body = slide.placeholders[1]
        _text(body.text_frame, [(level, text) for bullet in spec["bullets"]
                                for level, text in [(0, bullet["text"])] + [(1, sub) for sub in bullet["sub"]]])
        if spec["chart"] is not None:
            chart = charts[spec["chart"]]
            body.width = Inches(4.4)
            frame = _place_chart(slide, chart, draft_dir)
            frame._element.xpath("./*[1]/p:cNvPr")[0].set("descr", chart["fields"].get("Alt text", ""))
            given = {name: text for name, text in chart["fields"].items() if text and text.lower() != "none"}
            # A non-zero baseline's reason is in **Axis:**, and deck-rules.md keeps it in the source line.
            shown = ("Source", "Axis", "Interval") if non_zero_baseline(chart) else ("Source", "Interval")
            source = "; ".join(given[name] for name in shown if name in given)
            lines = [(0, text) for text in (source, given.get("Caveat") and f"* {given['Caveat']}") if text]
            if lines:
                source = slide.shapes.add_textbox(Inches(5.0), Inches(6.3), Inches(4.6), Inches(0.8))
                _text(source.text_frame, lines)
                for paragraph in source.text_frame.paragraphs:
                    paragraph.font.size = Pt(11)
        _notes(slide, spec["notes"])
    stream = io.BytesIO()
    deck.save(stream)
    output.write_bytes(_normalize(stream.getvalue()))
    return output


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__.rsplit("\n\n", 1)[-1])
    try:
        print(render(sys.argv[1], sys.argv[2]))
    except ContractError as error:
        sys.exit(f"refused: {error}")
